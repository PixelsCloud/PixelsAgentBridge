"""Generate deterministic document parser fixtures on Windows.

Requires PyMuPDF, Pillow, openpyxl and python-pptx. The output files are checked in so
Windows/macOS installed-client acceptance does not depend on Python packages.
"""
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

import fitz
from PIL import Image, ImageDraw, ImageFont
from openpyxl import Workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parent


def png_scan():
    image = Image.new("RGB", (1500, 420), "white")
    draw = ImageDraw.Draw(image)
    font = ImageFont.truetype("C:/Windows/Fonts/msyh.ttc", 62)
    draw.text((80, 70), "中文扫描合同 风险条款", font=font, fill="black")
    draw.text((80, 205), "Invoice total 12345 USD", font=font, fill="black")
    image.save(ROOT / "scan-zh-en.png")
    return ROOT / "scan-zh-en.png"


def pdfs(scan):
    native = fitz.open()
    page = native.new_page()
    page.insert_text((72, 90), "Pixels Risk Clause 7: payment due within 30 days.", fontsize=16)
    page.insert_text((72, 130), "Table: Item A | 42; Item B | 81", fontsize=13)
    native.save(ROOT / "native-table.pdf", deflate=True)
    native.close()

    scanned = fitz.open()
    page = scanned.new_page(width=750, height=210)
    page.insert_image(page.rect, filename=str(scan))
    scanned.save(ROOT / "scan-zh-en.pdf", deflate=True)
    scanned.close()

    mixed = fitz.open(ROOT / "native-table.pdf")
    mixed.insert_pdf(fitz.open(ROOT / "scan-zh-en.pdf"))
    mixed.save(ROOT / "mixed.pdf", deflate=True)
    mixed.close()

    locked = fitz.open(ROOT / "native-table.pdf")
    locked.save(
        ROOT / "encrypted.pdf", encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="fixture-owner", user_pw="fixture-user",
    )
    locked.close()


def spreadsheet():
    book = Workbook()
    first = book.active
    first.title = "Revenue"
    first["A1"] = "Item"
    first["B1"] = "Amount"
    first["A2"] = "Widget"
    first["B2"] = 42
    first["C2"] = "=B2*2"
    second = book.create_sheet("风险")
    second["B3"] = "逾期"
    second["C3"] = 81.5
    book.save(ROOT / "sheets-formula.xlsx")


def presentation():
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "风险汇报"
    first.placeholders[1].text = "项目 A：支付期限 30 天"
    first.notes_slide.notes_text_frame.text = "讲者备注：请复核第七条。"
    second = deck.slides.add_slide(deck.slide_layouts[1])
    second.shapes.title.text = "Summary"
    second.placeholders[1].text = "Item B has a value of 81."
    deck.save(ROOT / "slides-notes.pptx")


def compressed_bomb():
    # Small on disk, well above the parser's total expanded ZIP budget.
    with ZipFile(ROOT / "oversized-expanded.docx", "w", ZIP_DEFLATED) as archive:
        with archive.open("word/document.xml", "w", force_zip64=True) as output:
            block = b"0" * (1024 * 1024)
            for _ in range(129):
                output.write(block)


def ooxml():
    content_types = '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
 <Default Extension="xml" ContentType="application/xml"/>
 <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
 <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''
    body = '''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>
<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>风险评估</w:t></w:r></w:p>
<w:p><w:r><w:t>付款期限为三十天。</w:t></w:r></w:p>
<w:tbl><w:tr><w:tc><w:p><w:r><w:t>项目</w:t></w:r></w:p></w:tc>
<w:tc><w:p><w:r><w:t>金额</w:t></w:r></w:p></w:tc></w:tr>
<w:tr><w:tc><w:p><w:r><w:t>设备</w:t></w:r></w:p></w:tc>
<w:tc><w:p><w:r><w:t>42</w:t></w:r></w:p></w:tc></w:tr></w:tbl>
</w:body></w:document>'''
    with ZipFile(ROOT / "heading-table.docx", "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", '''<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>''')
        archive.writestr("word/document.xml", body)


if __name__ == "__main__":
    image = png_scan()
    pdfs(image)
    spreadsheet()
    presentation()
    compressed_bomb()
    ooxml()
    (ROOT / "malformed.pdf").write_bytes(b"%PDF-1.7\nthis is not a PDF")
